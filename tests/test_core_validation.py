import hashlib
import io
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from PIL import Image
from pypdf import PdfWriter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from review_core import parsing
from review_core.parsing import DocumentParseError, parse_document, validate_document_bytes
from review_core.redaction import REDACTION_POLICY_HASH, REDACTION_POLICY_VERSION, redact_sections
from review_core.reporting import FONT_PATH
from review_core.risk import RULESET_HASH, RULESET_VERSION, score_sections

MARKER = "SYNTHETIC DEMO DOCUMENT"


def _write(tmp_path: Path, content: bytes, name: str = "synthetic.txt") -> Path:
    path = tmp_path / name
    path.write_bytes(content)
    return path


def _pdf(*, image_page: bool = False, blank_page: bool = False) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(40, 760, MARKER)
    if image_page:
        pdf.showPage()
        from reportlab.lib.utils import ImageReader

        pdf.drawImage(ImageReader(Image.new("RGB", (20, 20), "white")), 40, 50, 200, 200)
    elif blank_page:
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("content", "name", "code"),
    [
        (b"", "synthetic.txt", "empty_document"),
        (b" \n\t", "synthetic.txt", "empty_text"),
        (b"\xff\xfe", "synthetic.txt", "invalid_encoding"),
        (b"fictional but unmarked", "synthetic.txt", "synthetic_marker_missing"),
        (b"not PDF", "synthetic.pdf", "type_mismatch"),
        (b"not ZIP", "synthetic.docx", "type_mismatch"),
        (b"%PDF-1.4\nnot valid", "synthetic.pdf", "invalid_pdf"),
        (MARKER.encode(), "synthetic.exe", "unsupported_type"),
        (b"%PDF-1.4\n", "synthetic.md", "type_mismatch"),
        ((MARKER + "\x00").encode(), "synthetic.txt", "unsupported_content"),
    ],
)
def test_invalid_input_is_not_a_successful_zero_risk_result(
    tmp_path: Path, content: bytes, name: str, code: str
) -> None:
    with pytest.raises(DocumentParseError) as exc:
        parse_document(_write(tmp_path, content, name))
    assert exc.value.code == code
    assert exc.value.retryable is False


def test_file_and_text_limits_are_utf8_byte_limits(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(parsing, "MAX_FILE_BYTES", 8)
    with pytest.raises(DocumentParseError, match="10 MiB") as exc:
        validate_document_bytes(b"x" * 9, ".txt")
    assert exc.value.code == "file_too_large"
    monkeypatch.setattr(parsing, "MAX_FILE_BYTES", 10 * 1024 * 1024)
    monkeypatch.setattr(parsing, "MAX_TEXT_BYTES", len(MARKER) + 4)
    with pytest.raises(DocumentParseError) as exc:
        parse_document(_write(tmp_path, (MARKER + "\n\n中文").encode()))
    assert exc.value.code == "text_too_large"


def test_wire_sections_are_stable_and_do_not_duplicate_full_text(tmp_path: Path) -> None:
    path = _write(tmp_path, (MARKER + "\n\n甲段。\n\n乙段。").encode())
    first, second = parse_document(path), parse_document(path)
    payload = first.to_dict()
    assert first == second
    assert "text" not in payload
    assert first.sha256 == hashlib.sha256(path.read_bytes()).hexdigest()
    assert [section.section_id for section in first.sections] == [
        "section-0001",
        "section-0002",
        "section-0003",
    ]
    assert [section.locator for section in first.sections] == ["block:1", "block:2", "block:3"]


def test_docx_tables_are_parsed_in_body_order_with_cell_locators(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph(MARKER)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "供应商承担全部损失。"
    table.cell(0, 1).text = "联系人：林晞，电话 13800138001。"
    document.add_paragraph("最后一段。")
    path = tmp_path / "synthetic.docx"
    document.save(path)
    parsed = parse_document(path)
    assert [section.locator for section in parsed.sections] == [
        "paragraph:1",
        "table:1/row:1/cell:1",
        "table:1/row:1/cell:2",
        "paragraph:2",
    ]
    redacted = redact_sections(parsed.to_dict()["sections"])
    result = score_sections(redacted.sections)
    legal = next(finding for finding in result.findings if finding.rule_id == "LEGAL-002")
    assert legal.locator == "table:1/row:1/cell:1"
    assert "13800138001" not in redacted.text


@pytest.mark.parametrize("structure", ["nested_table", "header", "textbox", "field_code"])
def test_unparsed_docx_structures_are_blocking(tmp_path: Path, structure: str) -> None:
    document = Document()
    document.add_paragraph(MARKER)
    if structure == "nested_table":
        document.add_table(rows=1, cols=1).cell(0, 0).add_table(rows=1, cols=1)
    elif structure == "header":
        document.sections[0].header.paragraphs[0].text = "隐藏责任条款。"
    elif structure == "textbox":
        document.paragraphs[0]._p.append(OxmlElement("w:txbxContent"))
    else:
        field = OxmlElement("w:instrText")
        field.text = "SYNTHETIC unsupported formula"
        document.paragraphs[0]._p.append(field)
    path = tmp_path / "synthetic.docx"
    document.save(path)
    with pytest.raises(DocumentParseError) as exc:
        parse_document(path)
    assert exc.value.code == "unsupported_content"


def test_docx_zip_boundary_and_resource_limits(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    document = Document()
    document.add_paragraph(MARKER)
    path = tmp_path / "synthetic.docx"
    document.save(path)
    valid_bytes = path.read_bytes()
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("../escape.txt", "SYNTHETIC")
    with pytest.raises(DocumentParseError) as exc:
        parse_document(path)
    assert exc.value.code == "invalid_docx"
    path.write_bytes(valid_bytes)
    monkeypatch.setattr(parsing, "MAX_DOCX_MEMBERS", 2)
    with pytest.raises(DocumentParseError) as exc:
        parse_document(path)
    assert exc.value.code == "docx_member_limit"
    monkeypatch.setattr(parsing, "MAX_DOCX_MEMBERS", 10_000)
    monkeypatch.setattr(parsing, "MAX_DOCX_EXPANDED_BYTES", 10)
    with pytest.raises(DocumentParseError) as exc:
        parse_document(path)
    assert exc.value.code == "docx_expanded_limit"


def test_non_word_zip_is_not_docx(tmp_path: Path) -> None:
    path = tmp_path / "synthetic.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("test.txt", MARKER)
    with pytest.raises(DocumentParseError) as exc:
        parse_document(path)
    assert exc.value.code == "type_mismatch"


def test_encrypted_and_image_only_pdf_are_explicit_failures(tmp_path: Path) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("synthetic-test-password")
    path = tmp_path / "encrypted.pdf"
    writer.write(path)
    with pytest.raises(DocumentParseError) as exc:
        parse_document(path)
    assert exc.value.code == "encrypted_pdf"
    with pytest.raises(DocumentParseError) as exc:
        parse_document(_write(tmp_path, _pdf(image_page=True), "synthetic.pdf"))
    assert exc.value.code == "requires_ocr"


def test_blank_pdf_page_is_informational_but_page_limit_is_blocking(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = _write(tmp_path, _pdf(blank_page=True), "synthetic.pdf")
    # ReportLab omits an untouched trailing page; append an explicit blank page.
    writer = PdfWriter(path)
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    path.write_bytes(buffer.getvalue())
    parsed = parse_document(path)
    assert any(warning.startswith("informational:blank_page:") for warning in parsed.warnings)
    monkeypatch.setattr(parsing, "MAX_PAGES", 1)
    with pytest.raises(DocumentParseError) as exc:
        parse_document(path)
    assert exc.value.code == "page_limit_exceeded"


def test_chinese_pdf_risk_has_real_page_locator_and_exact_evidence(tmp_path: Path) -> None:
    pdfmetrics.registerFont(TTFont("SyntheticInputCJK", str(FONT_PATH)))
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.setFont("SyntheticInputCJK", 11)
    pdf.drawString(40, 760, MARKER)
    pdf.showPage()
    pdf.setFont("SyntheticInputCJK", 11)
    pdf.drawString(40, 760, "供应商承担全部损失，联系人：林晞，13800138001。")
    pdf.save()
    parsed = parse_document(_write(tmp_path, buffer.getvalue(), "synthetic.pdf"))
    redacted = redact_sections(parsed.to_dict()["sections"])
    result = score_sections(redacted.sections)
    finding = next(item for item in result.findings if item.rule_id == "LEGAL-002")
    assert finding.locator == "page:2"
    assert redacted.sections[1]["text"][finding.start : finding.end] == finding.evidence
    assert "承担全部损失" in finding.evidence
    assert "13800138001" not in finding.evidence


def test_unicode_evidence_offsets_and_redaction_ids_are_document_wide() -> None:
    sections = [
        {"section_id": "s1", "locator": "paragraph:1", "text": "虚构😀联系人：林晞，13800138001。"},
        {
            "section_id": "s2",
            "locator": "paragraph:2",
            "text": "😀邮箱 other@example.org，供应商承担全部损失。",
        },
        {
            "section_id": "s3",
            "locator": "paragraph:3",
            "text": "再次共享账号，邮箱 next@example.org。",
        },
    ]
    redacted = redact_sections(sections)
    assert "[EMAIL_01]" in redacted.sections[1]["text"]
    assert "[EMAIL_02]" in redacted.sections[2]["text"]
    assessment = score_sections(redacted.sections)
    for finding in assessment.findings:
        section = next(
            item for item in redacted.sections if item["section_id"] == finding.section_id
        )
        assert section["text"][finding.start : finding.end] == finding.evidence
        assert finding.offset_basis == "unicode_code_points"
        assert finding.locator == section["locator"]
    assert assessment.to_dict()["ruleset_hash"] == RULESET_HASH
    assert assessment.ruleset_version == RULESET_VERSION
    assert redacted.to_dict()["policy_hash"] == REDACTION_POLICY_HASH
    assert redacted.policy_version == REDACTION_POLICY_VERSION
    assert score_sections(redacted.sections) == assessment


def test_repeated_rule_in_multiple_sections_preserves_unique_rule_scoring() -> None:
    result = score_sections(
        [
            {"section_id": "s1", "locator": "block:1", "text": "共享账号。"},
            {"section_id": "s2", "locator": "block:2", "text": "共享账号。"},
        ]
    )
    assert result.score == 26
    assert result.dimensions["security"] == 78
    assert len(result.findings) == 1
