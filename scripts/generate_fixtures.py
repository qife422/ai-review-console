from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "fixtures" / "documents"

HEADER = "SYNTHETIC DEMO DOCUMENT\nNorthstar Example Labs\n"
TEXT = (
    "负责人：林晞，邮箱 lin.xi@example.com，个人手机号 13800138001。\n"
    "本订单要求预付款90%，供应方承担全部损失，且允许使用共享账号。\n"
    "本文件完全由固定演示文本生成，不对应任何真实组织或合同。"
)


def main() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    (FIXTURES / "review-synthetic.txt").write_text(f"{HEADER}\n{TEXT}\n", encoding="utf-8")
    (FIXTURES / "review-synthetic.md").write_text(
        f"# {HEADER.splitlines()[0]}\n\n## Northstar Example Labs\n\n{TEXT}\n", encoding="utf-8"
    )

    document = Document()
    document.add_heading("SYNTHETIC DEMO DOCUMENT", level=1)
    document.add_paragraph("Northstar Example Labs")
    document.add_paragraph(TEXT)
    document.save(FIXTURES / "review-synthetic.docx")

    pdf = canvas.Canvas(str(FIXTURES / "review-synthetic.pdf"), pagesize=A4)
    pdf.setTitle("Synthetic Demo Document")
    y = A4[1] - 60
    for line in [*HEADER.splitlines(), "", *TEXT.splitlines()]:
        ascii_line = line.encode("ascii", "replace").decode("ascii")
        pdf.drawString(48, y, ascii_line)
        y -= 20
    pdf.save()
    print(f"Generated synthetic fixtures in {FIXTURES}")


if __name__ == "__main__":
    main()
