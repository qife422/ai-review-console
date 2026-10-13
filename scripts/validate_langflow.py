"""Exercise dynamic input, persistence, decisions and reports on a real stack.

Creates five explicitly synthetic cases, plus an optional bounded sequential
benchmark. No Flow is PATCHed and no existing document, decision or report is
overwritten.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import os
import platform
import subprocess
import sys
import time
import uuid
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import httpx

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from bootstrap_langflow import COMPONENTS, find_template  # noqa: E402

from api.app.langflow_client import LangflowClient, graph_hash  # noqa: E402

MIME_TYPES = {
    "txt": "text/plain",
    "md": "text/markdown",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def _items(payload: object) -> list[dict[str, Any]]:
    if isinstance(payload, dict):
        payload = payload.get("items", [])
    return [item for item in payload if isinstance(item, dict)] if isinstance(payload, list) else []


def _json(response: httpx.Response, expected: int = 200) -> Any:
    if response.status_code != expected:
        raise RuntimeError(
            f"Expected HTTP {expected}; got {response.status_code}: {response.text[:500]}"
        )
    return response.json()


def _wait_case(client: httpx.Client, case_id: str, *, reports: bool = False) -> dict[str, Any]:
    deadline = time.monotonic() + 160
    while time.monotonic() < deadline:
        detail = _json(client.get(f"/api/cases/{case_id}"))
        if reports:
            artifacts = detail.get("reports", [])
            if len(artifacts) >= 3 and all(item["state"] == "ready" for item in artifacts):
                return detail
            report_jobs = [job for job in detail.get("jobs", []) if job.get("kind") == "report"]
            latest = max(report_jobs, key=lambda job: job["created_at"]) if report_jobs else None
            if latest and latest["state"] == "failed":
                raise RuntimeError(f"Report job failed: {latest}")
        elif detail["status"] == "pending_approval":
            return detail
        elif detail["status"] == "failed":
            raise RuntimeError(f"Intake failed: {detail.get('jobs')}")
        time.sleep(1)
    raise TimeoutError(f"Case {case_id} did not reach the requested state")


def _format_bytes(extension: str, unique: str, clause: str) -> bytes:
    lines = ["SYNTHETIC DEMO DOCUMENT", "Built on Langflow", f"执行验证 {unique}", clause]
    if extension in {"txt", "md"}:
        return ("\n".join(lines[:3]) + "\n\n" + clause + "\n").encode()
    stream = io.BytesIO()
    if extension == "docx":
        from docx import Document

        document = Document()
        for line in lines[:3]:
            document.add_paragraph(line)
        document.add_table(rows=1, cols=1).cell(0, 0).text = clause
        document.save(stream)
    else:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen.canvas import Canvas

        font = Path(__file__).resolve().parents[1] / "review_core/assets/NotoSansSC-Regular.ttf"
        pdfmetrics.registerFont(TTFont("ARCValidationSC", str(font)))
        canvas = Canvas(stream, pagesize=(595, 842))
        canvas.setFont("ARCValidationSC", 12)
        for index, line in enumerate(lines):
            canvas.drawString(40, 780 - 26 * index, line)
        canvas.save()
    return stream.getvalue()


def _validate_case(
    client: httpx.Client, decision: str, index: int, extension: str
) -> dict[str, Any]:
    unique = uuid.uuid4().hex
    clauses = (
        "供应方承担全部损失，且不设责任上限。",
        "客户预付款90%，服务方可单方变更结算规则。",
        "联系人：林澈，个人手机号13800138001。使用共享账号和明文密码。",
    )
    content = _format_bytes(extension, unique, clauses[index])
    upload_key = f"flow-validation-upload-{unique}"
    upload_args = {
        "headers": {"Idempotency-Key": upload_key},
        "data": {"title": f"真实工作流虚构验证 {decision}", "synthetic_confirmed": "true"},
        "files": {
            "file": (f"flow-validation-{unique}.{extension}", content, MIME_TYPES[extension])
        },
    }
    accepted = _json(client.post("/api/cases", **upload_args), 202)
    case_id = accepted["case_id"]
    duplicate = _json(client.post("/api/cases", **upload_args), 202)
    if duplicate["case_id"] != case_id:
        raise AssertionError("Same-key upload created a duplicate case")
    detail = _wait_case(client, case_id)
    if (
        detail["source_kind"] != "uploaded_demo"
        or detail["document_hash"] != hashlib.sha256(content).hexdigest()
        or unique not in detail["original_text"]
        or not detail["findings"]
    ):
        raise AssertionError("Intake output was not bound to the newly uploaded bytes")
    intake_jobs = [item for item in detail["jobs"] if item["kind"] == "intake"]
    upstream_id = intake_jobs[-1].get("langflow_job_id") if intake_jobs else None
    if not upstream_id or str(upstream_id).startswith("lf-demo"):
        raise AssertionError("A real Langflow job ID is required")
    required_locator = "table:" if extension == "docx" else "page:" if extension == "pdf" else None
    if required_locator and not any(
        item["locator"].startswith(required_locator) for item in detail["findings"]
    ):
        raise AssertionError("The format-specific clause locator was not retained")
    reason = f"虚构验证完成：审查员依据规则证据记录{decision}，并核对中文报告。"
    decided = _json(
        client.post(
            f"/api/cases/{case_id}/decision",
            headers={
                "X-Reviewer": "synthetic-flow-validator",
                "Idempotency-Key": f"flow-validation-decision-{unique}",
            },
            json={
                "decision": decision,
                "reason": reason,
                "document_version_id": detail["current_document_version_id"],
                "assessment_id": detail["assessment_id"],
                "expected_row_version": detail["row_version"],
            },
        )
    )
    if decided["decision"] != decision:
        raise AssertionError("The human decision was not recorded")
    detail = _wait_case(client, case_id, reports=True)
    registered = []
    for artifact in detail["reports"]:
        report_format = artifact["format"]
        response = client.get(
            f"/api/cases/{case_id}/report/{report_format}",
            params={"report_id": artifact.get("report_id", artifact.get("id"))},
        )
        response.raise_for_status()
        digest = hashlib.sha256(response.content).hexdigest()
        if digest != artifact["sha256"] or len(response.content) != artifact["bytes"]:
            raise AssertionError("Downloaded artifact does not match registered bytes")
        if report_format == "json":
            snapshot = response.json()
            if (
                snapshot["case_id"] != case_id
                or snapshot["decision"] != decision
                or snapshot["decision_reason"] != reason
            ):
                raise AssertionError("Report used a different decision snapshot")
        registered.append(
            {"format": report_format, "sha256": digest, "bytes": len(response.content)}
        )
    report_jobs = [item for item in detail["jobs"] if item["kind"] == "report"]
    if not report_jobs[-1].get("langflow_job_id"):
        raise AssertionError("Report did not execute a real Decision Flow")
    print(f"Verified {extension.upper()} Intake / {decision} / three reports", file=sys.stderr)
    return {
        "case_id": case_id,
        "document_version_id": accepted["document_version_id"],
        "intake_job_id": accepted["job_id"],
        "intake_langflow_job_id": upstream_id,
        "decision": decision,
        "decision_id": detail["decision_id"],
        "report_langflow_job_id": report_jobs[-1]["langflow_job_id"],
        "source_sha256": hashlib.sha256(content).hexdigest(),
        "source_format": extension,
        "reports": registered,
    }


def _wait_failed(client: httpx.Client, case_id: str, job_id: str) -> dict[str, Any]:
    deadline = time.monotonic() + 160
    while time.monotonic() < deadline:
        detail = _json(client.get(f"/api/cases/{case_id}"))
        job = next(item for item in detail["jobs"] if item["job_id"] == job_id)
        if job["state"] == "failed" and detail["status"] == "failed":
            return detail
        if job["state"] == "succeeded":
            raise AssertionError("Missing synthetic marker was accepted by the real Flow")
        time.sleep(1)
    raise TimeoutError("Invalid input did not reach a reconciled terminal failure")


def _validate_failure_retry(client: httpx.Client) -> dict[str, Any]:
    unique = uuid.uuid4().hex
    # Deliberately omit the exact parser marker; the content remains wholly synthetic.
    content = f"专用虚构错误测试 {unique}\n仅用于演示，不包含真实资料。".encode()
    upload = {
        "headers": {"Idempotency-Key": f"invalid-input-upload-{unique}"},
        "data": {"title": "虚构缺标识失败验证", "synthetic_confirmed": "true"},
        "files": {"file": (f"invalid-marker-{unique}.txt", content, "text/plain")},
    }
    accepted = _json(client.post("/api/cases", **upload), 202)
    case_id, first_job_id = accepted["case_id"], accepted["job_id"]
    first = _wait_failed(client, case_id, first_job_id)
    first_job = next(item for item in first["jobs"] if item["job_id"] == first_job_id)
    if (
        first_job["error_code"] != "synthetic_marker_missing"
        or first_job["attempt"] != 1
        or first_job["retry_generation"] != 0
        or first_job["retryable"]
        or not first_job["langflow_job_id"]
        or any(
            first.get(field) is not None for field in ("risk_score", "risk_level", "assessment_id")
        )
    ):
        raise AssertionError("Invalid Intake did not expose a non-retryable unscored failure")
    retry_args = {
        "headers": {"Idempotency-Key": f"invalid-input-retry-{unique}"},
        "json": {
            "document_version_id": accepted["document_version_id"],
            "job_id": first_job_id,
        },
    }
    retry = _json(client.post(f"/api/cases/{case_id}/retry", **retry_args), 202)
    duplicate_retry = _json(client.post(f"/api/cases/{case_id}/retry", **retry_args), 202)
    if retry["job_id"] != duplicate_retry["job_id"] or retry["retry_generation"] != 1:
        raise AssertionError("Same-key manual retry created duplicate jobs")
    final = _wait_failed(client, case_id, retry["job_id"])
    duplicate_upload = _json(client.post("/api/cases", **upload), 202)
    events = [item["event_type"] for item in final["audit_events"]]
    retry_job = next(item for item in final["jobs"] if item["job_id"] == retry["job_id"])
    if (
        duplicate_upload["case_id"] != case_id
        or final["current_document_version_id"] != accepted["document_version_id"]
        or len(final["jobs"]) != 2
        or events.count("document_received") != 1
        or events.count("intake_retry_requested") != 1
        or events.count("intake_failed") != 2
        or any(item in events for item in ("case_scored", "decision_recorded", "report_ready"))
        or retry_job["error_code"] != "synthetic_marker_missing"
        or retry_job["attempt"] != 1
        or retry_job["retry_generation"] != 1
        or not retry_job["langflow_job_id"]
        or any(
            final.get(field) is not None for field in ("risk_score", "risk_level", "assessment_id")
        )
        or final["decision"] is not None
        or final["reports"]
    ):
        raise AssertionError("Manual retry duplicated or incorrectly scored business records")
    print("Verified missing-marker failure and deduplicated manual retry", file=sys.stderr)
    return {
        "case_id": case_id,
        "document_version_id": accepted["document_version_id"],
        "source_sha256": hashlib.sha256(content).hexdigest(),
        "first_job": first_job,
        "retry_job": retry_job,
        "same_key_upload_and_retry_deduplicated": True,
        "audit_event_counts": {name: events.count(name) for name in set(events)},
        "result": "failed as expected; unscored; same case/version; one manual retry generation",
    }


def _validate_markdown(client: httpx.Client) -> dict[str, Any]:
    unique = uuid.uuid4().hex
    content = _format_bytes("md", unique, "供应方承担全部损失且不设责任上限。")
    accepted = _json(
        client.post(
            "/api/cases",
            headers={"Idempotency-Key": f"markdown-validation-{unique}"},
            data={"title": "虚构 Markdown 类型验证", "synthetic_confirmed": "true"},
            files={"file": (f"format-probe-{unique}.md", content, "text/markdown")},
        ),
        202,
    )
    detail = _wait_case(client, accepted["case_id"])
    job = detail["jobs"][-1]
    if (
        detail["document_hash"] != hashlib.sha256(content).hexdigest()
        or unique not in detail["original_text"]
        or not detail["findings"]
        or not job["langflow_job_id"]
    ):
        raise AssertionError("Markdown Intake was not bound to its uploaded bytes")
    print("Verified Markdown Intake", file=sys.stderr)
    return {
        "case_id": accepted["case_id"],
        "source_format": "md",
        "source_sha256": hashlib.sha256(content).hexdigest(),
        "intake_job": job,
    }


def _benchmark(client: httpx.Client, count: int) -> dict[str, Any]:
    if not 1 <= count <= 20:
        raise ValueError("Benchmark count must be between 1 and 20")
    docker = subprocess.run(
        ["docker", "info", "--format", "{{json .}}"],
        check=True,
        capture_output=True,
        text=True,
        timeout=15,
    )
    info = json.loads(docker.stdout)
    samples, query_times = [], []
    for index in range(count):
        unique = uuid.uuid4().hex
        header = f"SYNTHETIC DEMO DOCUMENT\nBuilt on Langflow\n顺序基准 {unique}\n\n".encode()
        filler = "本条款仅为虚构演示，不包含真实信息。\n".encode()
        repeats, remainder = divmod(1024 * 1024 - len(header), len(filler))
        content = header + filler * repeats + b"z" * remainder
        start = time.monotonic()
        accepted = _json(
            client.post(
                "/api/cases",
                headers={"Idempotency-Key": f"sequential-benchmark-{unique}"},
                data={"title": f"虚构顺序基准 {index + 1}", "synthetic_confirmed": "true"},
                files={"file": (f"benchmark-{unique}.txt", content, "text/plain")},
            ),
            202,
        )
        upload_seconds = time.monotonic() - start
        detail = _wait_case(client, accepted["case_id"])
        intake_seconds = time.monotonic() - start
        job = detail["jobs"][-1]
        if (
            detail["document_hash"] != hashlib.sha256(content).hexdigest()
            or unique not in detail["original_text"]
            or job["state"] != "succeeded"
            or not job["langflow_job_id"]
        ):
            raise AssertionError("A sequential benchmark task was lost or used different bytes")
        for _ in range(3):
            query_start = time.monotonic()
            response = client.get(f"/api/cases/{accepted['case_id']}")
            response.raise_for_status()
            query_times.append(time.monotonic() - query_start)
        samples.append(
            {
                "case_id": accepted["case_id"],
                "job_id": accepted["job_id"],
                "langflow_job_id": job["langflow_job_id"],
                "source_bytes": len(content),
                "source_sha256": hashlib.sha256(content).hexdigest(),
                "upload_seconds": round(upload_seconds, 6),
                "intake_seconds": round(intake_seconds, 6),
            }
        )
        print(
            f"Benchmark {index + 1}/{count}: upload={upload_seconds:.3f}s "
            f"Intake={intake_seconds:.3f}s",
            file=sys.stderr,
            flush=True,
        )

    def p95(values: list[float]) -> float:
        return sorted(values)[math.ceil(len(values) * 0.95) - 1]

    upload_p95 = p95([item["upload_seconds"] for item in samples])
    query_p95 = p95(query_times)
    return {
        "environment": {
            "host_platform": platform.platform(),
            "host_machine": platform.machine(),
            "docker_server_version": info["ServerVersion"],
            "docker_cpu_count": info["NCPU"],
            "docker_memory_bytes": info["MemTotal"],
        },
        "method": "20-or-fewer sequential 1 MiB TXT uploads; loopback API; AI disabled",
        "samples": samples,
        "query_seconds": [round(item, 6) for item in query_times],
        "upload_p95_seconds": round(upload_p95, 6),
        "query_p95_seconds": round(query_p95, 6),
        "intake_max_seconds": max(item["intake_seconds"] for item in samples),
        "targets_met": {
            "upload_p95_le_3_seconds": upload_p95 <= 3,
            "query_p95_le_500_ms": query_p95 <= 0.5,
            "each_intake_le_120_seconds": all(item["intake_seconds"] <= 120 for item in samples),
            "sequential_tasks_not_lost": len(samples) == count,
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--api-url", default=os.getenv("REVIEW_API_URL", "http://localhost:8080"))
    parser.add_argument("--evidence", type=Path)
    parser.add_argument("--benchmark-count", type=int, default=0, choices=range(21))
    args = parser.parse_args()
    langflow = LangflowClient()
    try:
        component_index = langflow.request("GET", "/api/v1/all").json()
        missing = [name for name in COMPONENTS if not find_template(component_index, name)]
        if missing:
            raise RuntimeError(f"Missing trusted components: {', '.join(missing)}")
        flow_items = _items(langflow.request("GET", "/api/v1/flows/").json())
        selected = {
            name: next((item for item in flow_items if item.get("name") == name), None)
            for name in ("Document Review Intake", "Document Review Decision")
        }
        if not all(selected.values()):
            raise RuntimeError("Published review Flows are missing")
        before = {name: graph_hash(item["data"]) for name, item in selected.items() if item}
        openapi = langflow.request("GET", "/openapi.json").json()
        if openapi.get("info", {}).get("version") != "1.11.6":
            raise AssertionError("Runtime differs from the pinned Langflow 1.11.6 contract")
        build_operation = openapi["paths"]["/api/v1/build/{flow_id}/flow"]["post"]
        if not build_operation.get("requestBody"):
            raise AssertionError("Build API does not expose its per-execution input body")
        body_schema = build_operation["requestBody"]["content"]["application/json"]["schema"]
        schema_name = body_schema["$ref"].rsplit("/", 1)[-1]
        if "data" not in openapi["components"]["schemas"][schema_name]["properties"]:
            raise AssertionError("Build API does not expose a per-run data override")
        with httpx.Client(base_url=args.api_url.rstrip("/"), timeout=30) as api:
            results = [
                _validate_case(api, decision, index, extension)
                for index, (decision, extension) in enumerate(
                    (("approved", "txt"), ("rejected", "docx"), ("changes_requested", "pdf"))
                )
            ]
            markdown = _validate_markdown(api)
            failure_retry = _validate_failure_retry(api)
            benchmark = _benchmark(api, args.benchmark_count) if args.benchmark_count else None
        after = {
            name: graph_hash(langflow.request("GET", f"/api/v1/flows/{item['id']}").json()["data"])
            for name, item in selected.items()
            if item
        }
        if before != after:
            raise AssertionError("A per-case execution mutated a stored shared Flow")
        evidence = {
            "contract": "ARC-UPGRADE@0.1.0",
            "producer": "scripts/validate_langflow.py",
            "producer_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
            "command": sys.argv,
            "checked_at": datetime.now(UTC).isoformat(),
            "langflow_version": "1.11.6",
            "acceptance": ["AC-02", "AC-04", "AC-05", "AC-13", "AC-18"],
            "component_count": len(COMPONENTS),
            "build_input": "authenticated Build API data override; no per-case PATCH",
            "runtime_transport": "arc-runtime-v1: base64 of canonical UTF-8 JSON",
            "stored_flow_hashes_before": before,
            "stored_flow_hashes_after": after,
            "cases": results,
            "markdown_intake": markdown,
            "failure_retry": failure_retry,
            "benchmark": benchmark,
            "result": "passed",
        }
        if args.evidence:
            args.evidence.parent.mkdir(parents=True, exist_ok=True)
            args.evidence.write_text(
                json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8"
            )
        print(json.dumps(evidence, ensure_ascii=False, indent=2))
    finally:
        langflow.close()


if __name__ == "__main__":
    main()
